"""The shared hardened LibreOffice runner (B2b).

Conversion tests need a real LibreOffice; they skip when none is on PATH, and
run ungated against whatever version is present (the version pin is enforced
only by the worker entrypoint, never by convert). Timeout and crash use a fake
soffice script through NASHR_SOFFICE_BIN, so they are deterministic everywhere.
"""

from __future__ import annotations

import os
import shutil
import stat
from pathlib import Path

import pytest

from packages.platform import office_render
from packages.platform.office_render import (
    ConvertedFile,
    LibreOfficeUnavailableError,
    OfficeConversionError,
    OfficeTarget,
    OfficeTimeoutError,
    UnexpectedFontError,
    assert_libreoffice_version,
    convert,
)

_HAS_SOFFICE = shutil.which("soffice") is not None or shutil.which("libreoffice") is not None
_needs_soffice = pytest.mark.skipif(not _HAS_SOFFICE, reason="no LibreOffice on PATH")


def _make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for text in ("Birinchi slayd", "Ekinchi slayd"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        box.text_frame.text = text
    prs.save(str(path))


def _make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Maqola", level=1)
    document.add_paragraph("Bir paragraf matni.")
    document.save(str(path))


def _fake_soffice(path: Path, body: str) -> str:
    path.write_text(f"#!/usr/bin/env bash\n{body}\n", encoding="utf-8")
    path.chmod(path.stat().st_mode | stat.S_IEXEC | stat.S_IXGRP | stat.S_IXOTH)
    return str(path)


# ----------------------------------------------------------------- real renders


@_needs_soffice
async def test_pptx_converts_to_a_two_page_pdf(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    _make_pptx(src)
    result = await convert(src, OfficeTarget.PDF_FROM_PPTX, timeout_s=120)
    assert isinstance(result, ConvertedFile)
    assert result.pdf_bytes.startswith(b"%PDF-")
    assert result.page_count == 2


@_needs_soffice
async def test_docx_converts_to_pdf(tmp_path: Path) -> None:
    src = tmp_path / "article.docx"
    _make_docx(src)
    result = await convert(src, OfficeTarget.PDF_FROM_DOCX, timeout_s=120)
    assert result.pdf_bytes.startswith(b"%PDF-")
    assert result.page_count >= 1


@_needs_soffice
async def test_the_temp_dir_is_removed_after_a_conversion(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    _make_pptx(src)
    tmproot = Path(os.environ.get("TMPDIR", "/tmp"))
    before = set(tmproot.glob("nashr_office_*"))
    await convert(src, OfficeTarget.PDF_FROM_PPTX, timeout_s=120)
    after = set(tmproot.glob("nashr_office_*"))
    assert after <= before


@_needs_soffice
async def test_an_unexpected_font_is_refused(tmp_path: Path) -> None:
    src = tmp_path / "deck.pptx"
    _make_pptx(src)
    with pytest.raises(UnexpectedFontError):
        await convert(
            src,
            OfficeTarget.PDF_FROM_PPTX,
            timeout_s=120,
            allowed_font_families=frozenset({"NoSuchFamily"}),
        )


@_needs_soffice
async def test_concurrent_conversions_are_serialised(tmp_path: Path) -> None:
    import asyncio

    srcs = []
    for i in range(3):
        p = tmp_path / f"deck{i}.pptx"
        _make_pptx(p)
        srcs.append(p)
    results = await asyncio.gather(
        *(convert(p, OfficeTarget.PDF_FROM_PPTX, timeout_s=120) for p in srcs)
    )
    assert all(r.page_count == 2 for r in results)


# ----------------------------------------------------------------- fake soffice


async def test_a_timeout_kills_the_process_and_raises(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    fake = _fake_soffice(tmp_path / "soffice", "sleep 30")
    monkeypatch.setenv("NASHR_SOFFICE_BIN", fake)
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"not really a pptx")
    with pytest.raises(OfficeTimeoutError):
        await convert(src, OfficeTarget.PDF_FROM_PPTX, timeout_s=1)


async def test_a_nonzero_exit_raises_with_a_bounded_excerpt(
    tmp_path: Path,
    monkeypatch,  # type: ignore[no-untyped-def]
) -> None:
    fake = _fake_soffice(tmp_path / "soffice", "echo 'boom' >&2; exit 3")
    monkeypatch.setenv("NASHR_SOFFICE_BIN", fake)
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"x")
    with pytest.raises(OfficeConversionError) as info:
        await convert(src, OfficeTarget.PDF_FROM_PPTX, timeout_s=10)
    assert "boom" in str(info.value) and "exited 3" in str(info.value)


async def test_a_run_that_writes_no_pdf_raises(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    fake = _fake_soffice(tmp_path / "soffice", "exit 0")
    monkeypatch.setenv("NASHR_SOFFICE_BIN", fake)
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"x")
    with pytest.raises(OfficeConversionError):
        await convert(src, OfficeTarget.PDF_FROM_PPTX, timeout_s=10)


async def test_a_run_that_writes_an_unreadable_pdf_raises(
    tmp_path: Path,
    monkeypatch,  # type: ignore[no-untyped-def]
) -> None:
    # Write a non-PDF file at the expected output name.
    fake = _fake_soffice(
        tmp_path / "soffice",
        'outdir=""; prev=""; for a in "$@"; do if [ "$prev" = "--outdir" ]; then outdir="$a";'
        ' fi; prev="$a"; done; printf "garbage" > "$outdir/input.pdf"',
    )
    monkeypatch.setenv("NASHR_SOFFICE_BIN", fake)
    src = tmp_path / "deck.pptx"
    src.write_bytes(b"x")
    with pytest.raises(OfficeConversionError):
        await convert(src, OfficeTarget.PDF_FROM_PPTX, timeout_s=10)


def test_a_missing_binary_is_a_typed_error(monkeypatch) -> None:  # type: ignore[no-untyped-def]
    monkeypatch.setenv("NASHR_SOFFICE_BIN", "/nonexistent/soffice")
    with pytest.raises(LibreOfficeUnavailableError):
        office_render.find_soffice()


def test_the_version_gate_rejects_a_mismatch(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    fake = _fake_soffice(tmp_path / "soffice", "echo 'LibreOffice 1.2.3.4 abcdef'")
    monkeypatch.setenv("NASHR_SOFFICE_BIN", fake)
    with pytest.raises(LibreOfficeUnavailableError) as info:
        assert_libreoffice_version()
    assert office_render.PINNED_LIBREOFFICE_VERSION in str(info.value)


def test_the_version_gate_accepts_the_pin(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    pin = office_render.PINNED_LIBREOFFICE_VERSION
    fake = _fake_soffice(tmp_path / "soffice", f"echo 'LibreOffice {pin}.9.2 abcdef'")
    monkeypatch.setenv("NASHR_SOFFICE_BIN", fake)
    assert office_render.libreoffice_version().startswith(f"LibreOffice {pin}")
    assert assert_libreoffice_version().startswith("LibreOffice")
