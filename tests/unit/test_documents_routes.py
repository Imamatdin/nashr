"""Behavior tests for article document initialization, saves, history and export."""

from __future__ import annotations

from io import BytesIO
from typing import Any, cast
from uuid import UUID, uuid4

import httpx
import pytest
from docx import Document

from packages.api.app import create_app
from packages.api.services.tokens import mint_app_jwt
from packages.platform.config import PlatformConfig
from packages.platform.database import DatabaseClient
from packages.platform.document_models import DocumentKind, RevisionKind
from packages.platform.document_store import DocumentStore
from tests.unit.test_database_client import FakeSupabaseClient

pytestmark = pytest.mark.asyncio

_SECRET = "test-jwt-secret"
_USER_ID = uuid4()
_PROJECT_ID = str(uuid4())


def _content(title: str = "Learning") -> dict[str, Any]:
    return {"kind": "article", "schema_version": 1, "title": title, "blocks": []}


class _FakeDb:
    def __init__(self) -> None:
        self.owner = str(_USER_ID)
        self.project_type = "article"
        self.targets: list[dict[str, Any]] = []
        self.sources: list[dict[str, Any]] = []

    async def get_project_sources(self, project_id: str) -> list[dict[str, Any]]:
        return self.sources

    async def get_citation_targets(
        self, project_id: str, ids: list[str], *, claims: bool = False
    ) -> list[dict[str, Any]]:
        return [row for row in self.targets if row["id"] in ids]

    async def get_project(self, project_id: str) -> dict[str, Any] | None:
        if project_id != _PROJECT_ID:
            return None
        return {"id": project_id, "user_id": self.owner, "type": self.project_type}


def _config() -> PlatformConfig:
    return PlatformConfig(
        supabase_url="https://example.supabase.co",
        supabase_service_key="service",
        telegram_bot_token="123:abc",
        supabase_jwt_secret=_SECRET,
    )


class _Harness:
    def __init__(self) -> None:
        self.fake = FakeSupabaseClient()
        self.fake.seed(
            "projects", [{"id": _PROJECT_ID, "user_id": str(_USER_ID), "type": "article"}]
        )
        self.store = DocumentStore(DatabaseClient(_config(), client=cast(Any, self.fake)))
        self.db = _FakeDb()
        app = create_app(
            config=_config(),
            db=cast(Any, self.db),
            identity_service=cast(Any, object()),
            credits=cast(Any, object()),
            job_queue=cast(Any, object()),
            rate_limiter=cast(Any, object()),
            storage=cast(Any, object()),
            documents=self.store,
        )
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test"
        )

    async def seed(self, content: dict[str, Any] | None = None) -> str:
        """Give the article a head revision; returns its id."""

        actor = str(_USER_ID)
        document_id = await self.store.ensure_document(_PROJECT_ID, DocumentKind.ARTICLE, actor)
        head = await self.store.head(_PROJECT_ID, DocumentKind.ARTICLE)
        outcome = await self.store.commit(
            document_id,
            head.revision_id if head else None,
            content or _content(),
            kind=RevisionKind.MANUAL,
            actor=actor,
            idempotency_key=f"seed-{uuid4()}",
        )
        assert outcome.revision_id is not None
        return outcome.revision_id

    async def save(self, base: str, content: dict[str, Any], key: str) -> httpx.Response:
        return await self.client.put(
            f"/projects/{_PROJECT_ID}/documents/article",
            json={"base_revision_id": base, "idempotency_key": key, "content": content},
            headers=_headers(),
        )


def _headers(user_id: UUID = _USER_ID) -> dict[str, str]:
    session = mint_app_jwt(_SECRET, user_id, 3600)
    return {"Authorization": f"Bearer {session.access_token}"}


async def test_initialize_article_creates_empty_revision() -> None:
    h = _Harness()
    response = await h.client.post(
        f"/projects/{_PROJECT_ID}/documents/article",
        json={"title": "Learning", "idempotency_key": "init-1"},
        headers=_headers(),
    )
    assert response.status_code == 200
    assert response.json()["content"] == _content()
    head = await h.store.head(_PROJECT_ID, DocumentKind.ARTICLE)
    assert head is not None and head.revision_id == response.json()["revision_id"]


async def test_initialize_twice_returns_the_existing_head() -> None:
    h = _Harness()
    url = f"/projects/{_PROJECT_ID}/documents/article"
    first = await h.client.post(
        url, json={"title": "A", "idempotency_key": "i1"}, headers=_headers()
    )
    second = await h.client.post(
        url, json={"title": "B", "idempotency_key": "i2"}, headers=_headers()
    )
    assert second.json()["revision_id"] == first.json()["revision_id"]
    assert second.json()["content"]["title"] == "A"


async def test_save_article_advances_revision() -> None:
    h = _Harness()
    base = await h.seed()
    response = await h.save(base, _content("Changed"), "save-1")
    assert response.status_code == 200
    assert response.json()["revision_id"] != base
    head = await h.store.head(_PROJECT_ID, DocumentKind.ARTICLE)
    assert head is not None and head.content["title"] == "Changed"
    assert head.revision.parent_revision_id == base


async def test_passage_citation_is_checked_and_label_comes_from_source() -> None:
    h = _Harness()
    base = await h.seed()
    source_id, passage_id = str(uuid4()), str(uuid4())
    h.db.sources = [{"id": source_id, "filename": "study.pdf"}]
    h.db.targets = [{"id": passage_id, "source_id": source_id, "page": 3}]
    content = _content()
    content["blocks"] = [
        {
            "id": str(uuid4()),
            "type": "paragraph",
            "text": "My finding",
            "citations": [
                {"source_id": source_id, "passage_id": passage_id, "label": "Invented label"}
            ],
        }
    ]
    response = await h.save(base, content, "cite")
    assert response.status_code == 200
    citation = response.json()["content"]["blocks"][0]["citations"][0]
    assert citation["label"] == "study.pdf · 3"


@pytest.mark.parametrize("wrong_source", [False, True])
async def test_unknown_or_mismatched_citation_never_enters_a_revision(wrong_source: bool) -> None:
    h = _Harness()
    base = await h.seed()
    source_id, passage_id = str(uuid4()), str(uuid4())
    h.db.sources = [{"id": source_id, "filename": "study.pdf"}]
    if wrong_source:
        h.db.targets = [{"id": passage_id, "source_id": str(uuid4()), "page": 3}]
    content = _content()
    content["blocks"] = [
        {
            "id": str(uuid4()),
            "text": "Finding",
            "citations": [
                {"source_id": source_id, "passage_id": passage_id, "label": "Unverified"}
            ],
        }
    ]
    response = await h.save(base, content, "bad-cite")
    assert response.status_code == 422
    assert response.json()["detail"] == "citation_target_not_found"
    head = await h.store.head(_PROJECT_ID, DocumentKind.ARTICLE)
    assert head is not None and head.revision_id == base


async def test_stale_save_is_a_visible_conflict_that_names_the_head() -> None:
    h = _Harness()
    base = await h.seed()
    newer = (await h.save(base, _content("Other tab"), "tab-1")).json()["revision_id"]
    response = await h.save(base, _content("This tab"), "tab-2")
    assert response.status_code == 409
    assert response.json()["detail"] == {"reason": "stale_revision", "head_revision_id": newer}


async def test_other_user_cannot_read_document() -> None:
    h = _Harness()
    await h.seed()
    response = await h.client.get(
        f"/projects/{_PROJECT_ID}/documents/article", headers=_headers(uuid4())
    )
    assert response.status_code == 404


async def test_presentation_project_cannot_open_article_editor() -> None:
    h = _Harness()
    h.db.project_type = "presentation"
    response = await h.client.get(f"/projects/{_PROJECT_ID}/documents/article", headers=_headers())
    assert response.status_code == 409
    assert response.json()["detail"] == "not_article_project"


async def test_restore_commits_historical_content_as_new_head() -> None:
    h = _Harness()
    previous = await h.seed(_content("Previous"))
    current = await h.seed(_content("Current"))
    response = await h.client.post(
        f"/projects/{_PROJECT_ID}/documents/article/restore",
        json={"revision_id": previous, "base_revision_id": current, "idempotency_key": "r1"},
        headers=_headers(),
    )
    assert response.status_code == 200
    assert response.json()["content"]["title"] == "Previous"
    head = await h.store.head(_PROJECT_ID, DocumentKind.ARTICLE)
    assert head is not None and head.revision.kind is RevisionKind.RESTORE
    assert head.revision_id not in {previous, current}


async def test_export_article_returns_current_revision_as_docx() -> None:
    h = _Harness()
    head = await h.seed(
        {
            **_content("Evidence based teaching"),
            "blocks": [
                {
                    "id": str(uuid4()),
                    "type": "heading",
                    "text": "Methods",
                    "level": 2,
                    "citations": [],
                },
                {
                    "id": str(uuid4()),
                    "type": "paragraph",
                    "text": "Students compared two approaches.",
                    "level": None,
                    "citations": [],
                },
            ],
        }
    )
    response = await h.client.get(
        f"/projects/{_PROJECT_ID}/documents/article/export.docx", headers=_headers()
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in response.headers["content-disposition"]
    assert response.headers["x-document-revision"] == head
    rendered = Document(BytesIO(response.content))
    assert [paragraph.text for paragraph in rendered.paragraphs] == [
        "Evidence based teaching",
        "Methods",
        "Students compared two approaches.",
    ]


async def test_export_pins_the_visible_revision_when_another_tab_has_saved() -> None:
    """Downloading an older visible revision must not export the newer head."""
    h = _Harness()
    previous = await h.seed(_content("Previous"))
    await h.seed(_content("Newer content from another tab"))
    response = await h.client.get(
        f"/projects/{_PROJECT_ID}/documents/article/export.docx?revision_id={previous}",
        headers=_headers(),
    )
    assert response.status_code == 200
    assert response.headers["x-document-revision"] == previous
    rendered = Document(BytesIO(response.content))
    assert rendered.paragraphs[0].text == "Previous"


async def test_export_unknown_revision_does_not_fall_back_to_latest() -> None:
    """Missing historical content must produce a visible failure."""
    h = _Harness()
    await h.seed()
    response = await h.client.get(
        f"/projects/{_PROJECT_ID}/documents/article/export.docx?revision_id={uuid4()}",
        headers=_headers(),
    )
    assert response.status_code == 404
    assert response.json()["detail"] == "revision_not_found"


async def test_export_revision_requires_project_ownership() -> None:
    """A known revision id does not bypass the project's ownership boundary."""
    h = _Harness()
    previous = await h.seed()
    response = await h.client.get(
        f"/projects/{_PROJECT_ID}/documents/article/export.docx?revision_id={previous}",
        headers=_headers(uuid4()),
    )
    assert response.status_code == 404
