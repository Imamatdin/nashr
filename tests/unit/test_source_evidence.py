"""Extraction and reference validation against real local parsers."""

from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from pydantic import ValidationError
from reportlab.pdfgen.canvas import Canvas

from packages.core.models.document import DocumentCitation
from packages.workers.source.evidence import extract_passages


async def test_text_extraction_retains_content_without_model_calls() -> None:
    result = await extract_passages(
        b"Research evidence describes the observed result. " * 60, "research.txt"
    )
    assert len(result.chunks) > 1
    assert result.chunks[0].text.startswith("Research evidence")
    assert [chunk.page for chunk in result.chunks] == [None] * len(result.chunks)
    assert result.content_hash


async def test_docx_passages_do_not_claim_page_numbers() -> None:
    """DOCX parser 'pages' are heading sections; a page label would be invented."""

    document = Document()
    document.add_heading("Kirish", level=1)
    document.add_paragraph("Orol dengizi 1960-yildan boshlab qurib bormoqda.")
    document.add_heading("Oqibatlar", level=1)
    document.add_paragraph("Tuzli changlar aholining sog‘lig‘iga ta’sir ko‘rsatdi.")
    buffer = BytesIO()
    document.save(buffer)
    result = await extract_passages(buffer.getvalue(), "orol.docx")
    assert [chunk.page for chunk in result.chunks] == [None] * len(result.chunks)
    assert "Tuzli changlar" in " ".join(chunk.text for chunk in result.chunks)


async def test_pdf_extraction_preserves_page_locations() -> None:
    buffer = BytesIO()
    canvas = Canvas(buffer)
    canvas.drawString(50, 700, "First page: source evidence about learning.")
    canvas.showPage()
    canvas.drawString(50, 700, "Second page: independent observations.")
    canvas.save()
    result = await extract_passages(buffer.getvalue(), "study.pdf")
    assert [chunk.page for chunk in result.chunks] == [1, 2]
    assert "independent observations" in result.chunks[1].text


async def test_empty_upload_is_rejected() -> None:
    with pytest.raises(ValueError, match="invalid_source"):
        await extract_passages(b"", "empty.txt")


def test_passages_and_claims_are_distinct_citation_targets() -> None:
    source, passage = uuid4(), uuid4()
    citation = DocumentCitation(source_id=source, passage_id=passage, label="Study, p. 1")
    assert citation.claim_id is None
    assert citation.passage_id == passage
    assert DocumentCitation.model_validate(citation.model_dump()) == citation
    with pytest.raises(ValidationError, match="exactly one"):
        DocumentCitation(source_id=source, label="No target")
    with pytest.raises(ValidationError, match="exactly one"):
        DocumentCitation(source_id=source, passage_id=passage, claim_id=uuid4(), label="Ambiguous")
