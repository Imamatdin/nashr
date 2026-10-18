"""The one article builder: canonical blocks plus a profile become a DOCX.

This is the single place that turns an :class:`ArticleExportInput` into Word
bytes. Both the web route and the bot (through ``adapter.py``) reach the exporter
only through here. Layout, fonts, spacing and indent all come from the resolved
:class:`ArticleProfile`; citations are resolved against the input's ``sources``
map and rendered as in-text markers plus a bibliography.

python-docx is synchronous and exposes OOXML through unannotated lxml elements.
Every interaction with that untyped boundary is isolated in the ``_xml`` /
``_slot_*`` helpers below so the rest of the module stays strict-typed, the same
discipline the previous exporter used.
"""

from __future__ import annotations

import asyncio
import re
from io import BytesIO
from typing import Any, Final

from docx import Document  # pyright: ignore[reportUnknownVariableType]
from docx.document import Document as DocumentT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement  # pyright: ignore[reportUnknownVariableType]
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from packages.core.enums import CitationFormat, Language
from packages.core.models.bibliography import (
    CitationMetadata,
    FormattedBibliography,
    FormattedEntry,
)
from packages.core.models.document import ArticleBlock, ArticleBlockType
from packages.workers.article.bibliography import BibliographyFormatter
from packages.workers.article.export.labels import (
    ABSTRACT_HEADING,
    BIBLIOGRAPHY_HEADING,
    CHECKED_BY,
    KEYWORDS_LABEL,
    PERFORMED_BY,
    TOC_HEADING,
    article_type_label,
)
from packages.workers.article.export.models import (
    ArticleExportInput,
    BuiltArticle,
    UnresolvedCitationError,
)
from packages.workers.article.export.profiles import (
    Alignment,
    ArticleProfile,
    HeadingStyle,
    TitlePageKind,
    get_profile,
)

_SMALL_FONT_PT: Final[float] = 12.0
_HANGING_INDENT_CM: Final[float] = 1.25

_NUMBERED_BRACKET: Final[frozenset[CitationFormat]] = frozenset(
    {CitationFormat.GOST, CitationFormat.IEEE}
)
# Every style except APA references the reference list by number in this exporter.
_SUPERSCRIPT_DIGITS: Final[dict[str, str]] = {
    "0": "⁰",
    "1": "¹",
    "2": "²",
    "3": "³",
    "4": "⁴",
    "5": "⁵",
    "6": "⁶",
    "7": "⁷",
    "8": "⁸",
    "9": "⁹",
}

_ALIGN: Final[dict[Alignment, WD_ALIGN_PARAGRAPH]] = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


# ---------------------------------------------------------------------------
# OOXML boundary helpers (typed-Any by necessity; see module docstring)
# ---------------------------------------------------------------------------


def _xml(tag: str, **attrs: str) -> Any:
    elem: Any = OxmlElement(tag)  # pyright: ignore[reportUnknownVariableType]
    for key, value in attrs.items():
        elem.set(qn(key), value)  # pyright: ignore[reportUnknownMemberType]
    return elem  # pyright: ignore[reportUnknownVariableType]


def _append_run_xml(run: Any, *children: Any) -> None:
    for child in children:
        run._r.append(child)  # pyright: ignore[reportPrivateUsage]


def _set_xml_text(elem: Any, text: str) -> None:
    elem.text = text


def _section_pr(section: Any) -> Any:
    return section._sectPr  # pyright: ignore[reportPrivateUsage]


def _slot_fonts(style_element: Any, name: str) -> None:
    """Pin ascii, hAnsi, eastAsia and cs font slots so Cyrillic and Latin match.

    Without eastAsia and cs set, Word substitutes a default face for non-Latin
    glyphs even when the Latin name is set, producing mixed-typeface paragraphs.
    """

    rpr = style_element.find(qn("w:rPr"))
    if rpr is None:
        rpr = _xml("w:rPr")
        style_element.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = _xml("w:rFonts")
        rpr.append(rfonts)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(slot), name)


# ---------------------------------------------------------------------------
# Exporter
# ---------------------------------------------------------------------------


class ArticleExporter:
    """Builds a DOCX from an :class:`ArticleExportInput`. Stateless and reusable."""

    def __init__(self, formatter: BibliographyFormatter | None = None) -> None:
        self._formatter = formatter if formatter is not None else BibliographyFormatter()

    async def export(self, data: ArticleExportInput) -> BuiltArticle:
        """Build the DOCX off the event loop (python-docx is synchronous)."""

        return await asyncio.to_thread(self.build, data)

    def build(self, data: ArticleExportInput) -> BuiltArticle:
        """Build the DOCX synchronously and return it with its bibliography."""

        profile = get_profile(data.profile_id)
        lang = data.language
        doc = Document()
        _apply_page_layout(doc, profile)
        _apply_normal_style(doc, profile)
        _apply_heading_styles(doc, profile)
        _apply_page_numbers(doc, profile)

        bibliography, number_by_sid, cited_order = self._resolve_citations(data)

        if profile.title_page is not TitlePageKind.NONE:
            self._title_page(doc, data, profile)
            _page_break(doc)
        elif data.content.title:
            self._document_title(doc, data.content.title, profile)

        if data.include_toc:
            self._table_of_contents(doc, lang)
            _page_break(doc)

        if data.abstract:
            self._abstract(doc, data.abstract, profile, lang)
        if data.keywords:
            self._keywords(doc, data.keywords, profile, lang)

        word_count = 0
        citation_count = 0
        for block in data.content.blocks:
            used = self._render_block(
                doc, block, profile, data.citation_style, number_by_sid, data.sources, lang
            )
            citation_count += used
            word_count += len(block.text.split())

        if bibliography.entries:
            self._bibliography(doc, bibliography, profile, lang)

        buffer = BytesIO()
        doc.save(buffer)
        payload = buffer.getvalue()
        return BuiltArticle(
            docx_bytes=payload,
            filename=_filename(data),
            bibliography=bibliography,
            cited_source_ids=cited_order,
            word_count=word_count,
            block_count=len(data.content.blocks),
            citation_count=citation_count,
        )

    # -- citations ---------------------------------------------------------

    def _resolve_citations(
        self, data: ArticleExportInput
    ) -> tuple[FormattedBibliography, dict[str, int], list[str]]:
        """Collect cited sources in first-appearance order and build the bibliography.

        A citation whose source id is absent from ``data.sources`` raises
        :class:`UnresolvedCitationError`: we never emit a dangling marker.
        """

        appearance: list[str] = []
        seen: set[str] = set()
        for block in data.content.blocks:
            if block.type is ArticleBlockType.HEADING:
                # Headings render no in-text marker, so their citations must not
                # enter the reference list either; otherwise an entry would have
                # no marker pointing at it.
                continue
            for citation in block.citations:
                sid = str(citation.source_id)
                if sid not in data.sources:
                    raise UnresolvedCitationError(sid)
                if sid not in seen:
                    seen.add(sid)
                    appearance.append(sid)

        if not appearance:
            empty = FormattedBibliography(
                entries=[], style=data.citation_style, language=data.language.value, total_entries=0
            )
            return empty, {}, []

        ordered = [
            data.sources[sid].model_copy(update={"citation_number": index, "source_id": sid})
            for index, sid in enumerate(appearance, start=1)
        ]
        bibliography = self._build_bibliography(ordered, data.citation_style, data.language)
        number_by_sid = {
            entry.source_id: entry.number
            for entry in bibliography.entries
            if entry.source_id is not None and entry.number is not None
        }
        return bibliography, number_by_sid, appearance

    def _build_bibliography(
        self, ordered: list[CitationMetadata], style: CitationFormat, language: Language
    ) -> FormattedBibliography:
        """Build the reference list, numbered by appearance for every style but APA.

        APA is author-date and alphabetically ordered without numbers. Chicago is
        notes-bibliography by face but, with no footnote channel in the block
        model, we number it by appearance so its superscript markers resolve to a
        numbered reference list (recorded as a design decision in the lane file).
        """

        lang = language.value
        if style is CitationFormat.APA:
            return self._formatter.format_bibliography(ordered, style, lang)
        if style is CitationFormat.CHICAGO:
            entries = [
                FormattedEntry(
                    number=meta.citation_number,
                    formatted_text=self._formatter.format_single_reference(
                        meta, style, number=meta.citation_number, language=lang
                    ),
                    source_id=meta.source_id,
                    doi=meta.doi,
                )
                for meta in ordered
            ]
            return FormattedBibliography(
                entries=entries, style=style, language=lang, total_entries=len(entries)
            )
        return self._formatter.format_bibliography(ordered, style, lang)

    def _inline_marker(
        self,
        sids: list[str],
        style: CitationFormat,
        number_by_sid: dict[str, int],
        sources: dict[str, CitationMetadata],
        language: Language,
    ) -> str:
        """Render the in-text marker for one block's citations, grouped by style."""

        if style is CitationFormat.APA:
            parts = [
                self._formatter.format_inline_citation(
                    number_by_sid.get(sid, 1), sources[sid], style, language=language.value
                ).strip("()")
                for sid in sids
            ]
            return f"({'; '.join(parts)})"
        numbers = [number_by_sid[sid] for sid in sids]
        if style is CitationFormat.CHICAGO:
            return "".join(_superscript(n) for n in numbers)
        joined = ", ".join(str(n) for n in numbers)
        if style in _NUMBERED_BRACKET:
            return f"[{joined}]"
        return f"({joined})"  # Vancouver

    # -- blocks ------------------------------------------------------------

    def _render_block(
        self,
        doc: DocumentT,
        block: ArticleBlock,
        profile: ArticleProfile,
        style: CitationFormat,
        number_by_sid: dict[str, int],
        sources: dict[str, CitationMetadata],
        language: Language,
    ) -> int:
        """Render one canonical block; return how many citations it carried."""

        if block.type is ArticleBlockType.HEADING:
            self._heading(doc, block.text, block.level or 2, profile)
            # A heading carries no citations (model forbids body markers on it);
            # render any attached citations is out of scope, so count stays 0.
            return 0

        if block.type is ArticleBlockType.BULLET:
            paragraph = doc.add_paragraph(block.text, style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Cm(0)
        elif block.type is ArticleBlockType.QUOTE:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.left_indent = Cm(1)
            paragraph.paragraph_format.right_indent = Cm(1)
            quote_run: Any = paragraph.add_run(block.text)
            quote_run.italic = True
        else:  # paragraph
            paragraph = doc.add_paragraph(block.text)
            paragraph.style = "Normal"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if not block.citations:
            return 0

        sids: list[str] = []
        for citation in block.citations:
            sid = str(citation.source_id)
            if sid not in sids:
                sids.append(sid)
        marker = self._inline_marker(sids, style, number_by_sid, sources, language)
        # Superscript markers sit flush against the text; bracketed markers get a
        # leading space. Either way the marker is a real citation, never raw.
        if style is CitationFormat.CHICAGO:
            cite_run: Any = paragraph.add_run(marker)
            cite_run.font.superscript = True
        else:
            paragraph.add_run(f" {marker}")
        return len(sids)

    def _heading(self, doc: DocumentT, text: str, level: int, profile: ArticleProfile) -> None:
        style = profile.heading(level)
        paragraph = doc.add_paragraph(text.upper() if style.uppercase else text)
        paragraph.style = f"Heading {style.level}"
        paragraph.alignment = _ALIGN[style.alignment]

    def _document_title(self, doc: DocumentT, title: str, profile: ArticleProfile) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(18)
        run: Any = paragraph.add_run(title)
        run.bold = True
        run.font.name = profile.body_font
        run.font.size = Pt(profile.heading(1).size_pt + 2)

    # -- front matter ------------------------------------------------------

    def _table_of_contents(self, doc: DocumentT, language: Language) -> None:
        heading = doc.add_paragraph(TOC_HEADING[language])
        heading.style = "Heading 1"
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        begin = paragraph.add_run()
        _append_run_xml(begin, _xml("w:fldChar", **{"w:fldCharType": "begin"}))
        instr = paragraph.add_run()
        node = _xml("w:instrText", **{"xml:space": "preserve"})
        _set_xml_text(node, ' TOC \\o "1-3" \\h \\z \\u ')
        _append_run_xml(instr, node)
        separate = paragraph.add_run()
        _append_run_xml(separate, _xml("w:fldChar", **{"w:fldCharType": "separate"}))
        end = paragraph.add_run()
        _append_run_xml(end, _xml("w:fldChar", **{"w:fldCharType": "end"}))

    def _abstract(
        self, doc: DocumentT, text: str, profile: ArticleProfile, language: Language
    ) -> None:
        heading = doc.add_paragraph(ABSTRACT_HEADING[language])
        heading.style = "Heading 1"
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        body = doc.add_paragraph(text)
        body.style = "Normal"
        body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def _keywords(
        self, doc: DocumentT, keywords: list[str], profile: ArticleProfile, language: Language
    ) -> None:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run: Any = paragraph.add_run(f"{KEYWORDS_LABEL[language]}: {', '.join(keywords)}")
        run.italic = True
        run.font.name = profile.body_font
        run.font.size = Pt(_SMALL_FONT_PT)

    # -- title page --------------------------------------------------------

    def _title_page(
        self, doc: DocumentT, data: ArticleExportInput, profile: ArticleProfile
    ) -> None:
        if profile.title_page is TitlePageKind.STUDENT:
            self._student_title_page(doc, data, profile)
        else:
            self._apa_title_page(doc, data, profile)

    def _student_title_page(
        self, doc: DocumentT, data: ArticleExportInput, profile: ArticleProfile
    ) -> None:
        who = data.authorship
        if who.university_name:
            _centered(doc, who.university_name.upper(), profile, bold=True)
        if who.faculty_name:
            _centered(doc, who.faculty_name, profile)
        if who.department_name:
            _centered(doc, who.department_name, profile)
        _blank_lines(doc, 4)
        if data.article_type is not None:
            _centered(doc, article_type_label(data.article_type, data.language), profile, bold=True)
            _blank_lines(doc, 1)
        _centered(
            doc, data.content.title, profile, bold=True, size_pt=profile.heading(1).size_pt + 2
        )
        _blank_lines(doc, 4)
        if who.author_name:
            line = f"{PERFORMED_BY[data.language]}: {who.author_name}"
            if who.author_group:
                line += f", {who.author_group}"
            _aligned(doc, line, profile, WD_ALIGN_PARAGRAPH.RIGHT)
        if who.supervisor_name:
            line = f"{CHECKED_BY[data.language]}: {who.supervisor_name}"
            if who.supervisor_title:
                line += f", {who.supervisor_title}"
            _aligned(doc, line, profile, WD_ALIGN_PARAGRAPH.RIGHT)
        _blank_lines(doc, 6)
        locality = who.city or ""
        if who.year is not None:
            locality = f"{locality}, {who.year}" if locality else str(who.year)
        if locality:
            _centered(doc, locality, profile)

    def _apa_title_page(
        self, doc: DocumentT, data: ArticleExportInput, profile: ArticleProfile
    ) -> None:
        who = data.authorship
        _blank_lines(doc, 3)
        _centered(doc, data.content.title, profile, bold=True)
        _blank_lines(doc, 1)
        if who.author_name:
            _centered(doc, who.author_name, profile)
        affiliation = ", ".join(x for x in (who.department_name, who.university_name) if x)
        if affiliation:
            _centered(doc, affiliation, profile)
        if who.year is not None:
            _centered(doc, str(who.year), profile)

    # -- bibliography ------------------------------------------------------

    def _bibliography(
        self,
        doc: DocumentT,
        bibliography: FormattedBibliography,
        profile: ArticleProfile,
        language: Language,
    ) -> None:
        _page_break(doc)
        heading = doc.add_paragraph(BIBLIOGRAPHY_HEADING[language])
        heading.style = "Heading 1"
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for entry in bibliography.entries:
            paragraph = doc.add_paragraph()
            paragraph.style = "Normal"
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(-_HANGING_INDENT_CM)
            paragraph.paragraph_format.left_indent = Cm(_HANGING_INDENT_CM)
            paragraph.add_run(entry.formatted_text)


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------


def _apply_page_layout(doc: DocumentT, profile: ArticleProfile) -> None:
    section: Any = doc.sections[0]
    section.page_width = Cm(profile.page_width_cm)
    section.page_height = Cm(profile.page_height_cm)
    section.top_margin = Cm(profile.margin_top_cm)
    section.bottom_margin = Cm(profile.margin_bottom_cm)
    section.left_margin = Cm(profile.margin_left_cm)
    section.right_margin = Cm(profile.margin_right_cm)


def _apply_normal_style(doc: DocumentT, profile: ArticleProfile) -> None:
    style: Any = doc.styles["Normal"]  # pyright: ignore[reportUnknownMemberType]
    style.font.name = profile.body_font
    style.font.size = Pt(profile.body_size_pt)
    _slot_fonts(style.element, profile.body_font)
    pf = style.paragraph_format
    pf.line_spacing = profile.line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Cm(profile.first_line_indent_cm)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def _apply_heading_styles(doc: DocumentT, profile: ArticleProfile) -> None:
    for heading in profile.headings:
        _apply_one_heading(doc, profile, heading)


def _apply_one_heading(doc: DocumentT, profile: ArticleProfile, heading: HeadingStyle) -> None:
    style: Any = doc.styles[f"Heading {heading.level}"]  # pyright: ignore[reportUnknownMemberType]
    style.font.name = profile.body_font
    style.font.size = Pt(heading.size_pt)
    style.font.bold = heading.bold
    style.font.italic = heading.italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    _slot_fonts(style.element, profile.body_font)
    pf = style.paragraph_format
    pf.line_spacing = profile.line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0)
    pf.alignment = _ALIGN[heading.alignment]
    pf.space_before = Pt(heading.space_before_pt)
    pf.space_after = Pt(heading.space_after_pt)
    pf.keep_with_next = True


def _apply_page_numbers(doc: DocumentT, profile: ArticleProfile) -> None:
    if not profile.page_numbers.enabled:
        return
    section: Any = doc.sections[0]
    if profile.page_numbers.suppress_on_first_page:
        section.different_first_page_header_footer = True
        sect_pr = _section_pr(section)
        if sect_pr.find(qn("w:titlePg")) is None:
            sect_pr.append(_xml("w:titlePg"))
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = _ALIGN[profile.page_numbers.alignment]
    _append_page_field(paragraph)


def _append_page_field(paragraph: Any) -> None:
    begin = paragraph.add_run()
    _append_run_xml(begin, _xml("w:fldChar", **{"w:fldCharType": "begin"}))
    instr = paragraph.add_run()
    node = _xml("w:instrText", **{"xml:space": "preserve"})
    _set_xml_text(node, " PAGE ")
    _append_run_xml(instr, node)
    end = paragraph.add_run()
    _append_run_xml(end, _xml("w:fldChar", **{"w:fldCharType": "end"}))


# ---------------------------------------------------------------------------
# Small paragraph helpers
# ---------------------------------------------------------------------------


def _centered(
    doc: DocumentT,
    text: str,
    profile: ArticleProfile,
    *,
    bold: bool = False,
    size_pt: float | None = None,
) -> None:
    _aligned(doc, text, profile, WD_ALIGN_PARAGRAPH.CENTER, bold=bold, size_pt=size_pt)


def _aligned(
    doc: DocumentT,
    text: str,
    profile: ArticleProfile,
    alignment: WD_ALIGN_PARAGRAPH,
    *,
    bold: bool = False,
    size_pt: float | None = None,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run: Any = paragraph.add_run(text)
    run.bold = bold
    run.font.name = profile.body_font
    run.font.size = Pt(size_pt if size_pt is not None else profile.body_size_pt)


def _blank_lines(doc: DocumentT, count: int) -> None:
    for _ in range(count):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)


def _page_break(doc: DocumentT) -> None:
    paragraph = doc.add_paragraph()
    run: Any = paragraph.add_run()
    run.add_break()
    br = run._r.findall(qn("w:br"))[-1]  # pyright: ignore[reportPrivateUsage]
    br.set(qn("w:type"), "page")


def _superscript(number: int) -> str:
    return "".join(_SUPERSCRIPT_DIGITS[digit] for digit in str(number))


_FILENAME_STRIP: Final[re.Pattern[str]] = re.compile(r"[^\w\s-]", re.UNICODE)
_FILENAME_SPACES: Final[re.Pattern[str]] = re.compile(r"\s+")


def _filename(data: ArticleExportInput) -> str:
    safe = _FILENAME_STRIP.sub("", data.content.title)
    safe = _FILENAME_SPACES.sub("_", safe.strip())[:50].strip("_")
    if not safe:
        safe = "article"
    return f"{safe}.docx"


__all__ = ["ArticleExporter"]
